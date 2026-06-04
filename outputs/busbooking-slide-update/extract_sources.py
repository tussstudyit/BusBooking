from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def clean(text: str) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    return text


def extract_docx(path: Path) -> dict:
    doc = Document(str(path))
    paragraphs = []
    for p in doc.paragraphs:
        txt = clean(p.text)
        if not txt:
            continue
        style = p.style.name if p.style is not None else ""
        paragraphs.append({"style": style, "text": txt})

    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text) for cell in row.cells])
        if rows:
            tables.append(rows)

    return {"paragraphs": paragraphs, "tables": tables}


def extract_pptx_text(path: Path) -> list[dict]:
    slides = []
    with zipfile.ZipFile(path) as zf:
        names = sorted(
            [n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)],
            key=lambda n: int(re.search(r"slide(\d+)\.xml", n).group(1)),
        )
        for idx, name in enumerate(names, start=1):
            root = ET.fromstring(zf.read(name))
            runs = []
            for t in root.findall(".//a:t", NS):
                if t.text:
                    runs.append(t.text)
            text = clean(" ".join(runs))
            slides.append({"slide": idx, "text": text})
    return slides


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: extract_sources.py report.docx deck.pptx out.json")
    report = Path(sys.argv[1])
    deck = Path(sys.argv[2])
    out = Path(sys.argv[3])
    data = {
        "report": extract_docx(report),
        "existing_deck": extract_pptx_text(deck),
    }
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
